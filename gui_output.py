import streamlit as st
import pandas as pd
import os
import smtplib
import tempfile
import urllib.request
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from datetime import datetime, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from calendar import monthrange
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from supabase import create_client

FILE_DB = "database_monitoring_jtb.xlsx"
FILE_ALAT = "database_inspeksi_alat.xlsx"
FILE_TITIK = "master_titik_pantau.xlsx"
FILE_MASTER_ALAT = "master_daftar_alat.xlsx"
FILE_USER = "master_users.xlsx"
FILE_PENGATURAN = "pengaturan_laporan.xlsx"
FILE_KESIMPULAN = "master_kesimpulan.xlsx"
FILE_STATUS = "status_laporan.xlsx"
FILE_EMAIL_DIPERIKSA = "master_email_diperiksa.xlsx"
FILE_EMAIL_DISETUJUI = "master_email_disetujui.xlsx"
FILE_SMTP = "smtp_config.xlsx"
FOLDER_PETA = "peta_monitoring"
FOLDER_TTD = "tanda_tangan"

APP_LINK = "https://aplikasi-pemantauan-udara-hsse-pepc-zona-12-2ubay2uzebsaqvdkkc.streamlit.app/"

DEFAULT_TITIK = {
    "TP 1": {"Lokasi": "Pemukiman Ngasem", "Lat": "-7.1234", "Long": "111.6789", "LatNS": "-7.1234", "LongNS": "111.6789"},
    "TP 2": {"Lokasi": "Pemukiman Bandungrejo", "Lat": "-7.1245", "Long": "111.6791", "LatNS": "-7.1245", "LongNS": "111.6791"},
    "TP 3": {"Lokasi": "Petak 28", "Lat": "-7.1256", "Long": "111.6801", "LatNS": "-7.1256", "LongNS": "111.6801"},
    "TP 4": {"Lokasi": "Crossing Gayam", "Lat": "-7.1267", "Long": "111.6812", "LatNS": "-7.1267", "LongNS": "111.6812"},
    "TP 5": {"Lokasi": "Selatan Dormitory", "Lat": "-7.1278", "Long": "111.6823", "LatNS": "-7.1278", "LongNS": "111.6823"},
    "TP 6": {"Lokasi": "Dusun Besaran Ngasem", "Lat": "-7.1289", "Long": "111.6834", "LatNS": "-7.1289", "LongNS": "111.6834"},
}

DEFAULT_ALAT = ["Gas Detector 1", "Gas Detector 2", "SO2 Detector 1", "SO2 Detector 2"]

DEFAULT_KESIMPULAN = [
    "Tidak terdeteksi adanya Gas H2S dan SO2 di seluruh lokasi pemantauan",
    "Terdeteksi Gas H2S di bawah baku mutu",
    "Terdeteksi Gas SO2, perlu investigasi lanjutan",
    "Terdeteksi Gas H2S dan SO2, segera tindak lanjut",
]

STATUS_TEXT_COLS = [
    "Tanggal", "Status", "Diajukan_Oleh", "Waktu_Ajukan",
    "Diperiksa_Oleh", "Waktu_Diperiksa", "Disetujui_Oleh", "Waktu_Disetujui",
    "Ditolak_Oleh", "Waktu_Tolak", "Catatan_Tolak",
    "Email_Diperiksa", "Email_Disetujui",
]


def get_supabase():
    return create_client(st.secrets["supabase"]["url"], st.secrets["supabase"]["key"])


def teks_email_footer():
    return (
        "\n\nSilakan buka aplikasi untuk meninjau atau menyetujui laporan:\n"
        + APP_LINK
        + "\n\nHormat kami,\nSistem Pemantauan Udara Sesaat\nHSSE – PEPC Zona 12"
    )


def resolve_foto_path(path_str):
    if not path_str or str(path_str).strip() in ("", "nan", "None"):
        return None
    path_str = str(path_str).strip()

    if os.path.exists(path_str):
        return path_str

    local_try = os.path.join("foto_monitoring", path_str.replace("/", os.sep))
    if os.path.exists(local_try):
        return local_try

    if path_str.startswith("http://") or path_str.startswith("https://"):
        try:
            ext = os.path.splitext(path_str.split("?")[0])[1] or ".jpg"
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
            urllib.request.urlretrieve(path_str, tmp.name)
            return tmp.name
        except Exception:
            return None

    try:
        sb = get_supabase()
        pub = sb.storage.from_("foto-monitoring").get_public_url(path_str)
        if pub:
            ext = os.path.splitext(path_str)[1] or ".jpg"
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
            urllib.request.urlretrieve(pub, tmp.name)
            return tmp.name
    except Exception:
        pass
    return None


def load_monitoring_df():
    try:
        sb = get_supabase()
        res = sb.table("monitoring").select("*").order("id", desc=False).execute()
        rows = res.data or []
        if not rows:
            return pd.DataFrame(columns=[
                "Tanggal", "Bulan", "Shift", "Waktu", "Petugas", "Kode_Titik", "Lokasi",
                "H2S", "O2", "CO", "LEL", "SO2", "Noise", "Kesimpulan",
                "Path_Foto", "Path_Foto2", "Status", "id",
            ])
        df = pd.DataFrame(rows)
        rename = {
            "tanggal": "Tanggal", "bulan": "Bulan", "shift": "Shift", "waktu": "Waktu",
            "petugas": "Petugas", "kode_titik": "Kode_Titik", "lokasi": "Lokasi",
            "h2s": "H2S", "o2": "O2", "co": "CO", "lel": "LEL", "so2": "SO2",
            "noise": "Noise", "kesimpulan": "Kesimpulan",
            "path_foto": "Path_Foto", "path_foto2": "Path_Foto2", "status": "Status",
        }
        return df.rename(columns={k: v for k, v in rename.items() if k in df.columns})
    except Exception:
        if os.path.exists(FILE_DB):
            return pd.read_excel(FILE_DB)
        return pd.DataFrame()


def load_titik():
    try:
        sb = get_supabase()
        res = sb.table("titik_pantau").select("*").order("id", desc=False).execute()
        if res.data:
            titik = {}
            for row in res.data:
                kode = str(row.get("kode", "") or "")
                if not kode:
                    continue
                titik[kode] = {
                    "Lokasi": str(row.get("lokasi", "") or ""),
                    "Lat": str(row.get("latitude", "") or ""),
                    "Long": str(row.get("longitude", "") or ""),
                    "LatNS": str(row.get("latitude_ns", "") or row.get("latitude", "") or ""),
                    "LongNS": str(row.get("longitude_ns", "") or row.get("longitude", "") or ""),
                }
            if titik:
                return titik
    except Exception:
        pass
    if os.path.exists(FILE_TITIK):
        df = pd.read_excel(FILE_TITIK)
        titik = {}
        for _, row in df.iterrows():
            titik[str(row["Kode"])] = {
                "Lokasi": str(row["Lokasi"]),
                "Lat": str(row["Latitude"]),
                "Long": str(row["Longitude"]),
                "LatNS": str(row["Latitude_NS"]),
                "LongNS": str(row["Longitude_NS"]),
            }
        return titik
    return DEFAULT_TITIK


def load_titik_df():
    try:
        sb = get_supabase()
        res = sb.table("titik_pantau").select("*").order("id", desc=False).execute()
        if res.data:
            df = pd.DataFrame(res.data)
            rename = {
                "kode": "Kode", "lokasi": "Lokasi", "latitude": "Latitude",
                "longitude": "Longitude", "latitude_ns": "Latitude_NS", "longitude_ns": "Longitude_NS",
            }
            df = df.rename(columns={k: v for k, v in rename.items() if k in df.columns})
            cols = [c for c in ["id", "Kode", "Lokasi", "Latitude", "Longitude", "Latitude_NS", "Longitude_NS"] if c in df.columns]
            return df[cols]
    except Exception:
        pass
    if os.path.exists(FILE_TITIK):
        return pd.read_excel(FILE_TITIK)
    data = [{"Kode": k, "Lokasi": v["Lokasi"], "Latitude": v["Lat"], "Longitude": v["Long"],
             "Latitude_NS": v["LatNS"], "Longitude_NS": v["LongNS"]} for k, v in DEFAULT_TITIK.items()]
    return pd.DataFrame(data)


def load_daftar_alat():
    try:
        sb = get_supabase()
        res = sb.table("daftar_alat").select("nama_alat").order("id", desc=False).execute()
        if res.data:
            return [str(r["nama_alat"]) for r in res.data if r.get("nama_alat")]
    except Exception:
        pass
    if os.path.exists(FILE_MASTER_ALAT):
        return pd.read_excel(FILE_MASTER_ALAT)["Nama_Alat"].tolist()
    return list(DEFAULT_ALAT)


def load_kesimpulan_list():
    try:
        sb = get_supabase()
        res = sb.table("kesimpulan").select("teks").order("id", desc=False).execute()
        if res.data:
            return [str(r["teks"]) for r in res.data if r.get("teks")]
    except Exception:
        pass
    if os.path.exists(FILE_KESIMPULAN):
        df = pd.read_excel(FILE_KESIMPULAN)
        if "Kesimpulan" in df.columns:
            return df["Kesimpulan"].tolist()
    return list(DEFAULT_KESIMPULAN)


def load_pengaturan_laporan():
    default = {
        "Nama_Diperiksa": "Ian T.M",
        "Jabatan_Diperiksa": "Environment Officer",
        "Nama_Disetujui": "Ahmad Afifuddin",
        "Jabatan_Disetujui": "Superintendent Field HSSE",
    }
    try:
        sb = get_supabase()
        res = sb.table("pengaturan_laporan").select("*").order("id", desc=False).limit(1).execute()
        if res.data:
            r = res.data[0]
            return {
                "Nama_Diperiksa": str(r.get("nama_diperiksa") or default["Nama_Diperiksa"]),
                "Jabatan_Diperiksa": str(r.get("jabatan_diperiksa") or default["Jabatan_Diperiksa"]),
                "Nama_Disetujui": str(r.get("nama_disetujui") or default["Nama_Disetujui"]),
                "Jabatan_Disetujui": str(r.get("jabatan_disetujui") or default["Jabatan_Disetujui"]),
            }
    except Exception:
        pass
    if os.path.exists(FILE_PENGATURAN):
        try:
            df = pd.read_excel(FILE_PENGATURAN)
            return {
                "Nama_Diperiksa": str(df.loc[0, "Nama_Diperiksa"]) if "Nama_Diperiksa" in df.columns else default["Nama_Diperiksa"],
                "Jabatan_Diperiksa": str(df.loc[0, "Jabatan_Diperiksa"]) if "Jabatan_Diperiksa" in df.columns else default["Jabatan_Diperiksa"],
                "Nama_Disetujui": str(df.loc[0, "Nama_Disetujui"]) if "Nama_Disetujui" in df.columns else default["Nama_Disetujui"],
                "Jabatan_Disetujui": str(df.loc[0, "Jabatan_Disetujui"]) if "Jabatan_Disetujui" in df.columns else default["Jabatan_Disetujui"],
            }
        except Exception:
            return default
    return default


def _ensure_status_str(df):
    for c in STATUS_TEXT_COLS:
        if c not in df.columns:
            df[c] = ""
        df[c] = df[c].astype(object)
        df[c] = df[c].apply(lambda x: "" if pd.isna(x) or str(x).lower() == "nan" else str(x))
    return df


def load_status_laporan():
    try:
        sb = get_supabase()
        res = sb.table("status_laporan").select("*").order("id", desc=False).execute()
        rows = res.data or []
        if not rows:
            return pd.DataFrame(columns=STATUS_TEXT_COLS)
        df = pd.DataFrame(rows)
        rename = {
            "tanggal": "Tanggal", "status": "Status",
            "diajukan_oleh": "Diajukan_Oleh", "waktu_ajukan": "Waktu_Ajukan",
            "diperiksa_oleh": "Diperiksa_Oleh", "waktu_diperiksa": "Waktu_Diperiksa",
            "disetujui_oleh": "Disetujui_Oleh", "waktu_disetujui": "Waktu_Disetujui",
            "ditolak_oleh": "Ditolak_Oleh", "waktu_tolak": "Waktu_Tolak",
            "catatan_tolak": "Catatan_Tolak",
            "email_diperiksa": "Email_Diperiksa", "email_disetujui": "Email_Disetujui",
        }
        df = df.rename(columns={k: v for k, v in rename.items() if k in df.columns})
        return _ensure_status_str(df)
    except Exception:
        if os.path.exists(FILE_STATUS):
            return _ensure_status_str(pd.read_excel(FILE_STATUS))
        return pd.DataFrame(columns=STATUS_TEXT_COLS)


def get_status_tanggal(tgl_str):
    df = load_status_laporan()
    if df.empty:
        return "Draft"
    match = df[df["Tanggal"].astype(str) == str(tgl_str)]
    if len(match) == 0:
        return "Draft"
    return str(match.iloc[-1]["Status"])


def upsert_status(tgl_str, **kwargs):
    tgl_str = str(tgl_str)
    payload = {"tanggal": tgl_str}
    map_keys = {
        "Status": "status", "Diajukan_Oleh": "diajukan_oleh", "Waktu_Ajukan": "waktu_ajukan",
        "Diperiksa_Oleh": "diperiksa_oleh", "Waktu_Diperiksa": "waktu_diperiksa",
        "Disetujui_Oleh": "disetujui_oleh", "Waktu_Disetujui": "waktu_disetujui",
        "Ditolak_Oleh": "ditolak_oleh", "Waktu_Tolak": "waktu_tolak",
        "Catatan_Tolak": "catatan_tolak",
        "Email_Diperiksa": "email_diperiksa", "Email_Disetujui": "email_disetujui",
    }
    for k, v in kwargs.items():
        col = map_keys.get(k, k.lower())
        payload[col] = "" if v is None else str(v)
    try:
        sb = get_supabase()
        cek = sb.table("status_laporan").select("id").eq("tanggal", tgl_str).limit(1).execute()
        if cek.data:
            sb.table("status_laporan").update(payload).eq("id", cek.data[0]["id"]).execute()
        else:
            if "status" not in payload:
                payload["status"] = "Draft"
            sb.table("status_laporan").insert(payload).execute()
    except Exception as e:
        st.warning("Gagal simpan status Supabase: " + str(e))


def load_email_list(file_path, default_list=None):
    if default_list is None:
        default_list = []
    table = "email_diperiksa" if "diperiksa" in file_path else "email_disetujui"
    try:
        sb = get_supabase()
        res = sb.table(table).select("email").order("id", desc=False).execute()
        if res.data:
            return [str(r["email"]).strip() for r in res.data if r.get("email") and str(r["email"]).strip()]
    except Exception:
        pass
    if os.path.exists(file_path):
        df = pd.read_excel(file_path)
        if "Email" in df.columns:
            return [str(x).strip() for x in df["Email"].tolist() if str(x).strip() and str(x).lower() != "nan"]
    return list(default_list)


def load_smtp_config():
    try:
        em = str(st.secrets["smtp"]["email"]).strip()
        pw = str(st.secrets["smtp"]["password"]).strip().replace(" ", "")
        if em and pw:
            return em, pw
    except Exception:
        pass
    default_email = "pengirimaplikasi@gmail.com"
    default_pass = "spbckieaunuocqca"
    if os.path.exists(FILE_SMTP):
        try:
            df = pd.read_excel(FILE_SMTP)
            em = str(df.loc[0, "Email"]).strip() if "Email" in df.columns else default_email
            pw = str(df.loc[0, "Password"]).strip().replace(" ", "") if "Password" in df.columns else default_pass
            if not em or em.lower() == "nan":
                em = default_email
            if not pw or pw.lower() == "nan":
                pw = default_pass
            return em, pw
        except Exception:
            return default_email, default_pass
    return default_email, default_pass


def save_smtp_config(email, password):
    try:
        pd.DataFrame([{
            "Email": str(email).strip(),
            "Password": str(password).strip().replace(" ", ""),
        }]).to_excel(FILE_SMTP, index=False)
    except Exception:
        pass


def kirim_email(to_email, subject, body):
    to_email = str(to_email or "").strip()
    if not to_email or "@" not in to_email:
        return False, "Alamat email tujuan tidak valid"
    smtp_email, smtp_pass = load_smtp_config()
    try:
        msg = MIMEMultipart()
        msg["From"] = smtp_email
        msg["To"] = to_email
        msg["Subject"] = subject
        msg.attach(MIMEText(body, "plain", "utf-8"))
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=30) as server:
            server.login(smtp_email, smtp_pass)
            server.sendmail(smtp_email, [to_email], msg.as_string())
        return True, "Email terkirim ke " + to_email
    except Exception as e:
        return False, "Gagal kirim email: " + str(e)


def add_header_logo(canvas, doc):
    canvas.saveState()
    if os.path.exists("logo_skkmigas.png"):
        try:
            canvas.drawImage("logo_skkmigas.png", 1.8 * cm, A4[1] - 1.8 * cm, width=2.6 * cm, height=1.0 * cm, mask="auto")
        except Exception:
            pass
    if os.path.exists("logo_pertamina_epcepu.png"):
        try:
            canvas.drawImage("logo_pertamina_epcepu.png", A4[0] - 1.8 * cm - 3.0 * cm, A4[1] - 1.8 * cm, width=3.0 * cm, height=1.0 * cm, mask="auto")
        except Exception:
            pass
    canvas.restoreState()


def ambil_foto(df_shift, kode, tgl_str, shift_folder):
    path1, path2 = None, None
    match = df_shift[df_shift["Kode_Titik"] == kode] if not df_shift.empty else pd.DataFrame()
    if len(match) > 0:
        if "Path_Foto" in match.columns:
            p1 = match.iloc[0].get("Path_Foto", "")
            if p1 and str(p1).strip() not in ("", "nan", "None"):
                path1 = str(p1).strip()
        if "Path_Foto2" in match.columns:
            p2 = match.iloc[0].get("Path_Foto2", "")
            if p2 and str(p2).strip() not in ("", "nan", "None"):
                path2 = str(p2).strip()
    folder = "foto_monitoring/" + tgl_str + "/" + shift_folder
    if not path1:
        for ext in ["jpg", "jpeg", "png"]:
            for nomor in ["_1", ""]:
                candidate = folder + "/" + kode.replace(" ", "_") + nomor + "." + ext
                if os.path.exists(candidate):
                    path1 = candidate
                    break
            if path1:
                break
    if not path2:
        for ext in ["jpg", "jpeg", "png"]:
            candidate = folder + "/" + kode.replace(" ", "_") + "_2." + ext
            if os.path.exists(candidate):
                path2 = candidate
                break
    path1 = resolve_foto_path(path1) if path1 else None
    path2 = resolve_foto_path(path2) if path2 else None
    return path1, path2


def buat_excel_harian_resmi(df_db_all, tgl_target):
    titik_data = load_titik()
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Harian"
    font_title = Font(name="Arial", size=12, bold=True)
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="0288D1", end_color="0288D1", fill_type="solid")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A1:I1")
    ws.cell(row=1, column=1, value="LAPORAN PEMANTAUAN UDARA SESAAT - GPF JTB").font = font_title
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:I2")
    ws.cell(row=2, column=1, value="Tanggal: " + tgl_target.strftime("%d %B %Y") + " | Pertamina EP Cepu Zona 12").font = Font(name="Arial", size=10, italic=True)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=4, column=1, value="1. Daftar Titik Pantau dan Koordinat").font = font_title
    for col, text in enumerate(["No.", "Kode", "Lokasi", "Latitude", "Longitude", "Latitude NS", "Longitude NS"], 1):
        cell = ws.cell(row=5, column=col, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = center
        cell.border = thin_border
    for i, (kode, info) in enumerate(titik_data.items(), 1):
        row = 5 + i
        for col, val in enumerate([i, kode, info["Lokasi"], info["Lat"], info["Long"], info["LatNS"], info["LongNS"]], 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = thin_border
            cell.alignment = left if col == 3 else center
    df_f = pd.DataFrame()
    if not df_db_all.empty:
        try:
            df_f = df_db_all[df_db_all["Tanggal"].astype(str) == tgl_target.strftime("%Y-%m-%d")].reset_index(drop=True)
        except Exception:
            pass
    headers_ukur = ["No.", "Kode", "Lokasi", "Waktu", "H2S (ppm)", "O2 (%)", "CO (ppm)", "LEL (%)", "SO2 (ppm)"]
    start_day = 14
    ws.cell(row=start_day - 1, column=1, value="2. Hasil Pengukuran - Day Shift").font = font_title
    for col, text in enumerate(headers_ukur, 1):
        cell = ws.cell(row=start_day, column=col, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = center
        cell.border = thin_border
    df_day = df_f[df_f["Shift"] == "Day Shift"] if not df_f.empty else pd.DataFrame()
    for i, (kode, info) in enumerate(titik_data.items(), 1):
        row = start_day + i
        match = df_day[df_day["Kode_Titik"] == kode] if not df_day.empty else pd.DataFrame()
        if len(match) > 0:
            m = match.iloc[0]
            vals = [i, kode, info["Lokasi"], str(m.get("Waktu", "")),
                    "{:.2f}".format(float(m.get("H2S", 0) or 0)), "{:.1f}".format(float(m.get("O2", 0) or 0)),
                    "{:.1f}".format(float(m.get("CO", 0) or 0)), "{:.1f}".format(float(m.get("LEL", 0) or 0)),
                    "{:.2f}".format(float(m.get("SO2", 0) or 0))]
        else:
            vals = [i, kode, info["Lokasi"], "-", "-", "-", "-", "-", "-"]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = thin_border
            cell.alignment = left if col == 3 else center
    start_night = 23
    ws.cell(row=start_night - 1, column=1, value="3. Hasil Pengukuran - Night Shift").font = font_title
    for col, text in enumerate(headers_ukur, 1):
        cell = ws.cell(row=start_night, column=col, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = center
        cell.border = thin_border
    df_night = df_f[df_f["Shift"] == "Night Shift"] if not df_f.empty else pd.DataFrame()
    for i, (kode, info) in enumerate(titik_data.items(), 1):
        row = start_night + i
        match = df_night[df_night["Kode_Titik"] == kode] if not df_night.empty else pd.DataFrame()
        if len(match) > 0:
            m = match.iloc[0]
            vals = [i, kode, info["Lokasi"], str(m.get("Waktu", "")),
                    "{:.2f}".format(float(m.get("H2S", 0) or 0)), "{:.1f}".format(float(m.get("O2", 0) or 0)),
                    "{:.1f}".format(float(m.get("CO", 0) or 0)), "{:.1f}".format(float(m.get("LEL", 0) or 0)),
                    "{:.2f}".format(float(m.get("SO2", 0) or 0))]
        else:
            vals = [i, kode, info["Lokasi"], "-", "-", "-", "-", "-", "-"]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = thin_border
            cell.alignment = left if col == 3 else center
    ws.cell(row=32, column=1, value="Catatan: Berdasarkan hasil pemantauan udara sesaat menggunakan Multiple Gas Detector dan SO2 Detector.").font = Font(name="Arial", size=9, italic=True)
    for i, w in enumerate([6, 10, 22, 12, 12, 10, 12, 10, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def buat_excel_bulanan(df_db_all, bulan_target, tahun_target):
    bulan_map = {
        "Januari": 1, "Februari": 2, "Maret": 3, "April": 4, "Mei": 5, "Juni": 6,
        "Juli": 7, "Agustus": 8, "September": 9, "Oktober": 10, "November": 11, "Desember": 12,
    }
    titik_data = load_titik()
    wb = Workbook()
    ws = wb.active
    ws.title = bulan_target[:3] + "-" + str(tahun_target)[-2:]
    font_title = Font(name="Arial", size=14, bold=True)
    font_header = Font(name="Arial", size=8, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=7, bold=True)
    fill_header = PatternFill(start_color="0288D1", end_color="0288D1", fill_type="solid")
    fill_ds = PatternFill(start_color="BBDEFB", end_color="BBDEFB", fill_type="solid")
    fill_ns = PatternFill(start_color="FFE0B2", end_color="FFE0B2", fill_type="solid")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:Z1")
    ws.cell(row=1, column=1, value="Pemantauan Udara Sesaat").font = font_title
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    list_tp = list(titik_data.items())
    ws.cell(row=3, column=1, value=bulan_target).font = Font(name="Arial", size=10, bold=True)
    col = 2
    for kode, info in list_tp:
        start_col, end_col = col, col + 13
        ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
        cell = ws.cell(row=3, column=start_col, value=kode)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = center
        for c in range(start_col, end_col + 1):
            ws.cell(row=3, column=c).border = thin_border
            ws.cell(row=3, column=c).fill = fill_header
        col = end_col + 1
    col = 2
    for kode, info in list_tp:
        start_col, end_col = col, col + 13
        ws.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=end_col)
        cell = ws.cell(row=4, column=start_col, value=info["Lokasi"])
        cell.font = Font(name="Arial", size=8, bold=True)
        cell.alignment = center
        for c in range(start_col, end_col + 1):
            ws.cell(row=4, column=c).border = thin_border
        col = end_col + 1
    col = 2
    for _ in list_tp:
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 6)
        cell = ws.cell(row=5, column=col, value="Ds/")
        cell.font = font_sub
        cell.fill = fill_ds
        cell.alignment = center
        for c in range(col, col + 7):
            ws.cell(row=5, column=c).border = thin_border
            ws.cell(row=5, column=c).fill = fill_ds
        ws.merge_cells(start_row=5, start_column=col + 7, end_row=5, end_column=col + 13)
        cell = ws.cell(row=5, column=col + 7, value="Ns/")
        cell.font = font_sub
        cell.fill = fill_ns
        cell.alignment = center
        for c in range(col + 7, col + 14):
            ws.cell(row=5, column=c).border = thin_border
            ws.cell(row=5, column=c).fill = fill_ns
        col += 14
    params = ["Waktu", "H2S\n(ppm)", "O2\n(%)", "CO\n(ppm)", "LEL\n(%)", "SO2\n(ppm)", "Noise\ndB(A)"]
    col = 2
    for _ in list_tp:
        for p in params:
            cell = ws.cell(row=6, column=col, value=p)
            cell.font = Font(name="Arial", size=7, bold=True)
            cell.fill = fill_ds
            cell.alignment = center
            cell.border = thin_border
            col += 1
        for p in params:
            cell = ws.cell(row=6, column=col, value=p)
            cell.font = Font(name="Arial", size=7, bold=True)
            cell.fill = fill_ns
            cell.alignment = center
            cell.border = thin_border
            col += 1
    df_bulan = pd.DataFrame()
    if not df_db_all.empty:
        try:
            tmp = df_db_all.copy()
            tmp["Tanggal_dt"] = pd.to_datetime(tmp["Tanggal"], errors="coerce")
            bulan_angka = bulan_map.get(bulan_target, datetime.now().month)
            df_bulan = tmp[(tmp["Tanggal_dt"].dt.month == bulan_angka) & (tmp["Tanggal_dt"].dt.year == tahun_target)].copy()
        except Exception:
            pass
    bulan_angka = bulan_map.get(bulan_target, datetime.now().month)
    jumlah_hari = monthrange(tahun_target, bulan_angka)[1]
    for hari in range(1, jumlah_hari + 1):
        row = 6 + hari
        ws.cell(row=row, column=1, value=hari).border = thin_border
        ws.cell(row=row, column=1).alignment = center
        tgl_str = "{:04d}-{:02d}-{:02d}".format(tahun_target, bulan_angka, hari)
        col = 2
        for kode, info in list_tp:
            for shift, fill in [("Day Shift", fill_ds), ("Night Shift", fill_ns)]:
                match = pd.DataFrame()
                if not df_bulan.empty:
                    match = df_bulan[
                        (df_bulan["Tanggal"].astype(str) == tgl_str)
                        & (df_bulan["Kode_Titik"].astype(str) == kode)
                        & (df_bulan["Shift"] == shift)
                    ]
                if len(match) > 0:
                    m = match.iloc[0]
                    noise_val = ""
                    if pd.notna(m.get("Noise")):
                        try:
                            noise_val = "{:.1f}".format(float(m.get("Noise", 0)))
                        except Exception:
                            noise_val = ""
                    vals = [
                        str(m.get("Waktu", "")),
                        "{:.1f}".format(float(m.get("H2S", 0) or 0)),
                        "{:.1f}".format(float(m.get("O2", 0) or 0)),
                        "{:.1f}".format(float(m.get("CO", 0) or 0)),
                        "{:.1f}".format(float(m.get("LEL", 0) or 0)),
                        "{:.1f}".format(float(m.get("SO2", 0) or 0)),
                        noise_val,
                    ]
                else:
                    vals = ["", "0", "0", "0", "0", "0", ""]
                for v in vals:
                    cell = ws.cell(row=row, column=col, value=v)
                    cell.border = thin_border
                    cell.alignment = center
                    cell.fill = fill
                    col += 1
    note_row = 6 + jumlah_hari + 2
    ws.cell(row=note_row, column=1, value="Note : Berdasarkan hasil pemantauan udara sesaat dengan menggunakan Multiple Gas Detector dan SO2 Detector di beberapa titik.").font = Font(name="Arial", size=8, italic=True)
    ws.column_dimensions["A"].width = 5
    for i in range(2, 2 + len(list_tp) * 14):
        ws.column_dimensions[get_column_letter(i)].width = 6
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
def buat_pdf_laporan(df_db, tgl_target):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=1.8 * cm, leftMargin=1.8 * cm,
        topMargin=2.5 * cm, bottomMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    style_title = ParagraphStyle("TitleCustom", parent=styles["Heading1"], fontSize=12, alignment=TA_CENTER, spaceAfter=2, fontName="Helvetica-Bold")
    style_subtitle = ParagraphStyle("Subtitle", parent=styles["Normal"], fontSize=11, alignment=TA_CENTER, spaceAfter=3, fontName="Helvetica-Bold")
    style_heading = ParagraphStyle("HeadingCustom", parent=styles["Heading2"], fontSize=11, spaceBefore=5, spaceAfter=2, fontName="Helvetica-Bold")
    style_normal = ParagraphStyle("NormalCustom", parent=styles["Normal"], fontSize=10, leading=12, alignment=TA_JUSTIFY, fontName="Helvetica")
    style_small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=8, leading=10, alignment=TA_LEFT, fontName="Helvetica")
    style_center = ParagraphStyle("Center", parent=styles["Normal"], fontSize=10, alignment=TA_CENTER)
    style_left = ParagraphStyle("Left", parent=styles["Normal"], fontSize=10, alignment=TA_LEFT)
    style_caption = ParagraphStyle("Caption", parent=styles["Normal"], fontSize=9, leading=11, alignment=TA_CENTER, fontName="Helvetica")

    story = []
    titik_data = load_titik()
    tgl_str = tgl_target.strftime("%Y-%m-%d")

    story.append(Spacer(1, 3))
    story.append(Paragraph("LAPORAN PEMANTAUAN UDARA SESAAT", style_title))
    story.append(Paragraph("PADA KEGIATAN PRODUKSI FASE OPERASI", style_subtitle))
    story.append(Spacer(1, 2))
    story.append(Paragraph(
        "Berikut disampaikan laporan kegiatan pemantauan sesaat dengan menggunakan "
        "<i>Multiple Gas Detector</i> dan SO2 detector di Area GPF dan sekitar GPF "
        "pada saat kegiatan Produksi Fase Operasi.",
        style_normal,
    ))
    story.append(Spacer(1, 3))
    story.append(Paragraph("1. Pelaksanaan :", style_heading))

    df_hari = pd.DataFrame()
    if not df_db.empty and "Tanggal" in df_db.columns:
        df_hari = df_db[df_db["Tanggal"].astype(str) == tgl_str].copy()

    petugas_ds, petugas_ns = "-", "-"
    jam_ds_awal, jam_ds_akhir = "-", "-"
    jam_ns_awal, jam_ns_akhir = "-", "-"

    if not df_hari.empty:
        ds = df_hari[df_hari["Shift"] == "Day Shift"] if "Shift" in df_hari.columns else pd.DataFrame()
        ns = df_hari[df_hari["Shift"] == "Night Shift"] if "Shift" in df_hari.columns else pd.DataFrame()
        if not ds.empty:
            petugas_ds = str(ds.iloc[0].get("Petugas", "-"))
            waktu_ds = sorted(ds["Waktu"].astype(str).tolist()) if "Waktu" in ds.columns else []
            if waktu_ds:
                jam_ds_awal, jam_ds_akhir = waktu_ds[0], waktu_ds[-1]
        if not ns.empty:
            petugas_ns = str(ns.iloc[0].get("Petugas", "-"))
            waktu_ns = sorted(ns["Waktu"].astype(str).tolist()) if "Waktu" in ns.columns else []
            if waktu_ns:
                jam_ns_awal, jam_ns_akhir = waktu_ns[0], waktu_ns[-1]

    hari_list = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
    bulan_list = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
                  "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    tgl_label = (
        hari_list[tgl_target.weekday()] + ", " + str(tgl_target.day) + " "
        + bulan_list[tgl_target.month] + " " + str(tgl_target.year)
    )

    pel_data = [
        ["Hari, Tanggal", ":", tgl_label],
        ["Jam", ":", "Day Shift " + jam_ds_awal + " WIB s/d " + jam_ds_akhir + " WIB"],
        ["", "", "Night Shift " + jam_ns_awal + " WIB s/d " + jam_ns_akhir + " WIB"],
        ["Tempat", ":", "Luar Area GPF"],
        ["Dilaporkan oleh", ":", ""],
        ["Petugas Sampling", ":", "Environment (DS) & ERCM (NS)"],
        ["Siang dilakukan pengukuran oleh", ":", petugas_ds],
        ["Malam dilakukan pengukuran oleh", ":", petugas_ns],
    ]
    t_pel = Table(pel_data, colWidths=[5.5 * cm, 0.4 * cm, 10.5 * cm])
    t_pel.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    story.append(t_pel)

    story.append(Paragraph("2. Tujuan :", style_heading))
    story.append(Paragraph("a. Memastikan dispersi dari kegiatan flaring JTB GPF aman.", style_normal))
    story.append(Paragraph("b. Sebagai salah satu data rona lingkungan pada Tahap Operasi berlangsung.", style_normal))

    story.append(Paragraph("3. Titik Pengukuran", style_heading))
    story.append(Paragraph(
        "Titik monitoring dilakukan di luar area GPF, penentuan titik pemantauan berdasarkan 4 arah mata angin serta kawasan permukiman terdekat:",
        style_normal,
    ))
    story.append(Spacer(1, 2))

    data_titik = [["No.", "Kode", "Lokasi", "Latitude", "Longitude", "Latitude NS", "Longitude NS"]]
    for i, (kode, info) in enumerate(titik_data.items(), 1):
        data_titik.append([str(i), kode, info["Lokasi"], info["Lat"], info["Long"], info["LatNS"], info["LongNS"]])
    t_titik = Table(data_titik, colWidths=[1.0 * cm, 1.5 * cm, 4.0 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm])
    t_titik.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0288D1")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(t_titik)

    story.append(Paragraph("4. Hasil Pengukuran", style_heading))
    story.append(Paragraph(
        "Hasil monitoring pada titik pantau menggunakan <i>Multiple Gas Detector</i> adalah sebagai berikut:",
        style_normal,
    ))

    def ada_noise_df(dfx):
        if dfx is None or dfx.empty or "Noise" not in dfx.columns:
            return False
        for v in dfx["Noise"].tolist():
            try:
                if v is not None and str(v).strip() not in ("", "nan", "None") and float(v) > 0:
                    return True
            except Exception:
                pass
        return False

    pakai_noise = ada_noise_df(df_hari)

    def buat_tabel_hasil(df_shift, judul):
        elems = [Paragraph("<b>" + judul + "</b>", style_small)]
        if pakai_noise:
            data = [["No.", "Kode", "Lokasi", "Waktu", "H2S (ppm)", "O2 (%)", "CO (ppm)", "LEL (%)", "SO2 (ppm)", "Kebisingan (dB)"]]
            widths = [0.7 * cm, 1.2 * cm, 3.0 * cm, 1.6 * cm, 1.5 * cm, 1.3 * cm, 1.4 * cm, 1.3 * cm, 1.5 * cm, 1.8 * cm]
        else:
            data = [["No.", "Kode", "Lokasi", "Waktu", "H2S (ppm)", "O2 (%)", "CO (ppm)", "LEL (%)", "SO2 (ppm)"]]
            widths = [0.8 * cm, 1.3 * cm, 3.5 * cm, 1.8 * cm, 1.8 * cm, 1.5 * cm, 1.6 * cm, 1.5 * cm, 1.7 * cm]

        for i, (kode, info) in enumerate(titik_data.items(), 1):
            match = df_shift[df_shift["Kode_Titik"] == kode] if (not df_shift.empty and "Kode_Titik" in df_shift.columns) else pd.DataFrame()
            if len(match) > 0:
                m = match.iloc[0]
                row = [
                    str(i), kode, info["Lokasi"], str(m.get("Waktu", "")),
                    "{:.1f}".format(float(m.get("H2S", 0) or 0)),
                    "{:.1f}".format(float(m.get("O2", 0) or 0)),
                    "{:.1f}".format(float(m.get("CO", 0) or 0)),
                    "{:.1f}".format(float(m.get("LEL", 0) or 0)),
                    "{:.1f}".format(float(m.get("SO2", 0) or 0)),
                ]
                if pakai_noise:
                    try:
                        nv = m.get("Noise", None)
                        if nv is not None and str(nv).strip() not in ("", "nan", "None"):
                            row.append("{:.1f}".format(float(nv)))
                        else:
                            row.append("-")
                    except Exception:
                        row.append("-")
            else:
                row = [str(i), kode, info["Lokasi"], "-", "0", "0", "0", "0", "0"]
                if pakai_noise:
                    row.append("-")
            data.append(row)

        t = Table(data, colWidths=widths)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0288D1")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7 if pakai_noise else 8),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        elems.append(t)
        return elems

    df_day = df_hari[df_hari["Shift"] == "Day Shift"] if (not df_hari.empty and "Shift" in df_hari.columns) else pd.DataFrame()
    df_night = df_hari[df_hari["Shift"] == "Night Shift"] if (not df_hari.empty and "Shift" in df_hari.columns) else pd.DataFrame()
    story.append(Spacer(1, 2))
    story.extend(buat_tabel_hasil(df_day, "Day Shift"))
    story.append(Spacer(1, 3))
    story.extend(buat_tabel_hasil(df_night, "Night Shift"))

    def buat_blok_foto(caption, path1, path2):
        ada1 = path1 and os.path.exists(path1)
        ada2 = path2 and os.path.exists(path2)
        if ada1 and ada2:
            try:
                img1 = Image(path1, width=3.95 * cm, height=5.29 * cm)
            except Exception:
                img1 = Paragraph("[Foto error]", style_center)
            try:
                img2 = Image(path2, width=3.95 * cm, height=5.29 * cm)
            except Exception:
                img2 = Paragraph("[Foto error]", style_center)
            inner = Table([[img1, img2]], colWidths=[4.3 * cm, 4.3 * cm])
            inner.setStyle(TableStyle([
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]))
            isi_foto = inner
        elif ada1 or ada2:
            p = path1 if ada1 else path2
            try:
                isi_foto = Image(p, width=3.95 * cm, height=5.29 * cm)
            except Exception:
                isi_foto = Paragraph("[Foto error]", style_center)
        else:
            isi_foto = Paragraph("[Belum ada foto]", style_center)
        t = Table([[Paragraph(caption, style_caption)], [isi_foto]], colWidths=[16.4 * cm])
        t.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.8, colors.black),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F8F8F8")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    story.append(PageBreak())
    story.append(Paragraph("5. Dokumentasi, Pemantauan Day Shift", style_heading))
    for kode, info in titik_data.items():
        match = df_day[df_day["Kode_Titik"] == kode] if (not df_day.empty and "Kode_Titik" in df_day.columns) else pd.DataFrame()
        waktu = str(match.iloc[0]["Waktu"]) if len(match) > 0 else "..."
        caption = "Pemantauan di " + kode + " (" + info["Lokasi"] + ") pukul " + waktu + " WIB"
        path1, path2 = ambil_foto(df_day, kode, tgl_str, "Day_Shift")
        story.append(buat_blok_foto(caption, path1, path2))
        story.append(Spacer(1, 3))

    story.append(PageBreak())
    path_peta1 = resolve_foto_path("peta/peta_1.jpg") or os.path.join(FOLDER_PETA, "peta_1.jpg")
    if path_peta1 and os.path.exists(path_peta1):
        try:
            img_peta = Image(path_peta1, width=15.22 * cm, height=7.41 * cm)
            img_peta.hAlign = "CENTER"
            story.append(img_peta)
            story.append(Spacer(1, 4))
        except Exception:
            pass

    story.append(Paragraph("6. Dokumentasi, pemantauan Night Shift", style_heading))
    story.append(Spacer(1, 2))
    for kode, info in titik_data.items():
        match = df_night[df_night["Kode_Titik"] == kode] if (not df_night.empty and "Kode_Titik" in df_night.columns) else pd.DataFrame()
        waktu = str(match.iloc[0]["Waktu"]) if len(match) > 0 else "..."
        caption = "Pemantauan di " + kode + " (" + info["Lokasi"] + ") pukul " + waktu + " WIB"
        path1, path2 = ambil_foto(df_night, kode, tgl_str, "Night_Shift")
        story.append(buat_blok_foto(caption, path1, path2))
        story.append(Spacer(1, 3))

    path_peta2 = resolve_foto_path("peta/peta_2.jpg") or os.path.join(FOLDER_PETA, "peta_2.jpg")
    if path_peta2 and os.path.exists(path_peta2):
        try:
            img_peta2 = Image(path_peta2, width=15.22 * cm, height=7.41 * cm)
            img_peta2.hAlign = "CENTER"
            story.append(img_peta2)
        except Exception:
            pass

    story.append(Spacer(1, 8))
    story.append(Paragraph("7. Kesimpulan :", style_heading))
    kesimpulan_hari = "-"
    if not df_hari.empty and "Kesimpulan" in df_hari.columns:
        kes_list = [x for x in df_hari["Kesimpulan"].dropna().tolist() if str(x).strip()]
        if kes_list:
            kesimpulan_hari = str(kes_list[-1])
    story.append(Paragraph(kesimpulan_hari, style_normal))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "Bojonegoro, " + str(tgl_target.day) + " " + bulan_list[tgl_target.month] + " " + str(tgl_target.year),
        style_left,
    ))
    story.append(Spacer(1, 6))

    peng = load_pengaturan_laporan()
    status_lap = get_status_tanggal(tgl_str)
    tampil_dip = status_lap in ["Menunggu Disetujui", "Final"]
    tampil_dis = status_lap == "Final"

    def get_ttd_img(path):
        if os.path.exists(path):
            try:
                return Image(path, width=3.5 * cm, height=2.2 * cm)
            except Exception:
                return ""
        base = os.path.basename(path)
        resolved = resolve_foto_path("ttd/" + base)
        if resolved and os.path.exists(resolved):
            try:
                return Image(resolved, width=3.5 * cm, height=2.2 * cm)
            except Exception:
                return ""
        return ""

    img_ds = get_ttd_img(os.path.join(FOLDER_TTD, "ttd_day.png"))
    img_ns = get_ttd_img(os.path.join(FOLDER_TTD, "ttd_night.png"))
    img_dip = get_ttd_img(os.path.join(FOLDER_TTD, "ttd_diperiksa.png")) if tampil_dip else ""
    img_dis = get_ttd_img(os.path.join(FOLDER_TTD, "ttd_disetujui.png")) if tampil_dis else ""

    ttd_data = [
        [Paragraph("<b>Dilaporkan oleh</b>", style_center), "",
         Paragraph("<b>Diperiksa oleh</b>", style_center),
         Paragraph("<b>Disetujui Oleh</b>", style_center)],
        [Paragraph("<b>Day Shift</b>", style_center),
         Paragraph("<b>Night Shift</b>", style_center), "", ""],
        [img_ds if img_ds else "", img_ns if img_ns else "",
         img_dip if img_dip else "", img_dis if img_dis else ""],
        [
            Paragraph("<b>" + petugas_ds + "</b><br/><b>Associate Environment Staff</b>", style_center),
            Paragraph("<b>" + petugas_ns + "</b><br/><b>ERCM Staff</b>", style_center),
            Paragraph("<b>" + peng["Nama_Diperiksa"] + "</b><br/><b>" + peng["Jabatan_Diperiksa"] + "</b>", style_center),
            Paragraph("<b>" + peng["Nama_Disetujui"] + "</b><br/><b>" + peng["Jabatan_Disetujui"] + "</b>", style_center),
        ],
    ]
    t_ttd = Table(ttd_data, colWidths=[4.2 * cm, 4.2 * cm, 4.2 * cm, 4.2 * cm])
    t_ttd.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
        ("LINEAFTER", (0, 0), (0, -1), 0.5, colors.black),
        ("LINEAFTER", (1, 0), (1, -1), 0.5, colors.black),
        ("LINEAFTER", (2, 0), (2, -1), 0.5, colors.black),
        ("SPAN", (0, 0), (1, 0)),
        ("LINEBELOW", (0, 1), (1, 1), 0, colors.white),
        ("LINEABOVE", (0, 3), (-1, 3), 0, colors.white),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_ttd)

    story.append(Spacer(1, 10))
    story.append(Paragraph("<i>Keterangan:</i>", style_small))
    story.append(Paragraph("1. Nilai Ambang Batas <b>H2S</b> sebesar <b>1 ppm</b> mengacu pada Peraturan Menteri Ketenagakerjaan Nomor 5 Tahun 2018 Tentang Keselamatan dan Kesehatan Kerja Lingkungan Kerja", style_small))
    story.append(Paragraph("2. Rentang nilai ketersediaan <b>oksigen</b> di udara sebesar <b>19,5% – 23,5%</b> mengacu pada Peraturan Menteri Ketenagakerjaan Nomor 5 Tahun 2018 Tentang Keselamatan dan Kesehatan Kerja Lingkungan Kerja", style_small))
    story.append(Paragraph("3. Baku Mutu <b>CO</b> sebesar <b>8,72 ppm</b> mengacu pada Peraturan Pemerintah (PP) Nomor 22 Tahun 2021 tentang Penyelenggaraan Perlindungan dan Pengelolaan Lingkungan Hidup", style_small))
    story.append(Paragraph("4. Baku Mutu <b>SO2</b> sebesar <b>0,057 ppm</b> mengacu pada Peraturan Pemerintah (PP) Nomor 22 Tahun 2021 tentang Penyelenggaraan Perlindungan dan Pengelolaan Lingkungan Hidup", style_small))
    story.append(Paragraph("5. Baku Tingkat <b>Kebisingan</b> sebesar <b>55 dB</b> untuk Kawasan Perumahan dan Pemukiman mengacu pada Keputusan Menteri Negara Lingkungan Hidup Nomor 48 Tahun 1996", style_small))

    doc.build(story, onFirstPage=add_header_logo, onLaterPages=add_header_logo)
    buffer.seek(0)
    return buffer.getvalue()
def tampilkan_halaman_output(pilihan_menu, nama_user="", username="", role="petugas"):
    titik_data = load_titik()
    daftar_alat = load_daftar_alat()

    if pilihan_menu == "2. Review Data":
        st.markdown("### 2. Portal Pencarian & Review Data")
        df_db = load_monitoring_df()
        if df_db.empty:
            st.info("Belum ada data pemantauan.")
            return
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            cari_tgl = st.date_input("Tanggal", datetime.now())
        with c2:
            cari_shift = st.selectbox("Shift", ["Semua", "Day Shift", "Night Shift"])
        with c3:
            cari_tp = st.selectbox("Titik", ["Semua"] + list(titik_data.keys()))
        with c4:
            cari_status = st.selectbox("Status", ["Semua", "Aman", "Waspada", "Bahaya"])
        df_hasil = df_db[df_db["Tanggal"].astype(str) == cari_tgl.strftime("%Y-%m-%d")].copy()
        if cari_shift != "Semua" and "Shift" in df_hasil.columns:
            df_hasil = df_hasil[df_hasil["Shift"] == cari_shift]
        if cari_tp != "Semua" and "Kode_Titik" in df_hasil.columns:
            df_hasil = df_hasil[df_hasil["Kode_Titik"].astype(str) == cari_tp]
        if cari_status != "Semua" and "Status" in df_hasil.columns:
            df_hasil = df_hasil[df_hasil["Status"] == cari_status]
        if not df_hasil.empty:
            if "Status" in df_hasil.columns:
                if "Bahaya" in df_hasil["Status"].values:
                    st.error("TERDETEKSI STATUS BAHAYA")
                elif "Waspada" in df_hasil["Status"].values:
                    st.warning("ADA STATUS WASPADA")
                else:
                    st.success("SEMUA TITIK AMAN")

            def style_row(row):
                s = str(row.Status) if "Status" in row.index else ""
                if s in ["Bahaya", "Terpapar"]:
                    return ["background-color:#FFEBEE"] * len(row)
                if s in ["Waspada", "O2 Tidak Normal"]:
                    return ["background-color:#FFF8E1"] * len(row)
                return ["background-color:#E8F5E9"] * len(row)

            tampil = df_hasil.drop(columns=["id", "created_at"], errors="ignore")
            try:
                st.dataframe(tampil.style.apply(style_row, axis=1), use_container_width=True)
            except Exception:
                st.dataframe(tampil, use_container_width=True)
            st.markdown("---")
            if "id" in df_hasil.columns and not df_hasil["id"].isna().all():
                opsi = df_hasil.index.tolist()
                idx = st.selectbox(
                    "Pilih baris untuk dihapus:",
                    opsi,
                    format_func=lambda i: (
                        str(df_hasil.loc[i].get("Kode_Titik", "")) + " | "
                        + str(df_hasil.loc[i].get("Shift", "")) + " | "
                        + str(df_hasil.loc[i].get("Waktu", "")) + " | id="
                        + str(df_hasil.loc[i].get("id", ""))
                    ),
                )
                if st.button("Hapus Permanen"):
                    try:
                        rid = int(df_hasil.loc[idx]["id"])
                        get_supabase().table("monitoring").delete().eq("id", rid).execute()
                        st.success("Data dihapus dari Supabase")
                        st.rerun()
                    except Exception as e:
                        st.error("Gagal hapus: " + str(e))
        else:
            st.info("Belum ada data pada filter ini.")

    elif pilihan_menu == "3. Download Laporan":
        st.markdown("### 3. Panel Cetak & Download")
        df_db = load_monitoring_df()
        tab1, tab2, tab3 = st.tabs(["Excel Harian", "Excel Bulanan", "PDF Laporan Resmi"])
        with tab1:
            st.caption("Unduh Excel Harian:")
            tgl = st.date_input("Pilih Tanggal", datetime.now(), key="tgl_harian")
            if st.button("Generate Excel Harian", type="primary"):
                data = buat_excel_harian_resmi(df_db, tgl)
                st.download_button(
                    "Download", data=data,
                    file_name="Laporan_Harian_" + tgl.strftime("%Y%m%d") + ".xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        with tab2:
            st.caption("Unduh Excel Bulanan:")
            c1, c2 = st.columns(2)
            with c1:
                bln = st.selectbox(
                    "Bulan",
                    ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
                     "Juli", "Agustus", "September", "Oktober", "November", "Desember"],
                    index=datetime.now().month - 1,
                )
            with c2:
                thn = st.number_input("Tahun", 2024, 2030, datetime.now().year)
            if st.button("Generate Excel Bulanan", type="primary"):
                data = buat_excel_bulanan(df_db, bln, thn)
                st.download_button(
                    "Download", data=data,
                    file_name="Laporan_Bulanan_" + bln + "_" + str(thn) + ".xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        with tab3:
            st.caption("PDF: petugas setelah Final. Officer/Superintendent bisa preview.")
            tgl_pdf = st.date_input("Pilih Tanggal untuk PDF", datetime.now(), key="tgl_pdf")
            tgl_pdf_str = tgl_pdf.strftime("%Y-%m-%d")
            status_now = get_status_tanggal(tgl_pdf_str)
            warna = {
                "Draft": "#64748B", "Menunggu Diperiksa": "#F9A825",
                "Menunggu Disetujui": "#1565C0", "Final": "#2E7D32", "Ditolak": "#C62828",
            }
            st.markdown(
                "**Status:** <span style='color:" + warna.get(status_now, "#333")
                + ";font-weight:700;'>" + status_now + "</span>",
                unsafe_allow_html=True,
            )
            df_st = load_status_laporan()
            row_st = df_st[df_st["Tanggal"].astype(str) == tgl_pdf_str] if not df_st.empty else pd.DataFrame()
            if len(row_st) > 0 and status_now == "Ditolak":
                cat = str(row_st.iloc[-1].get("Catatan_Tolak", "") or "")
                if cat.strip():
                    st.error("Alasan ditolak: " + cat)
            df_hari_pdf = df_db[df_db["Tanggal"].astype(str) == tgl_pdf_str].copy() if not df_db.empty else pd.DataFrame()
            list_tp = list(titik_data.keys())
            total_tp = len(list_tp) if list_tp else 6

            def hitung_shift(df, shift):
                if df.empty or "Shift" not in df.columns:
                    return 0, list_tp
                d = df[df["Shift"] == shift]
                sudah = d["Kode_Titik"].astype(str).unique().tolist() if not d.empty and "Kode_Titik" in d.columns else []
                return len(sudah), [t for t in list_tp if t not in sudah]

            n_ds, belum_ds = hitung_shift(df_hari_pdf, "Day Shift")
            n_ns, belum_ns = hitung_shift(df_hari_pdf, "Night Shift")
            ada_kes = False
            if not df_hari_pdf.empty and "Kesimpulan" in df_hari_pdf.columns:
                ada_kes = any(str(x).strip() and str(x) != "nan" for x in df_hari_pdf["Kesimpulan"].dropna())
            jml_foto = 0
            if not df_hari_pdf.empty:
                for col_f in ["Path_Foto", "Path_Foto2"]:
                    if col_f in df_hari_pdf.columns:
                        jml_foto += int(
                            df_hari_pdf[col_f].fillna("").astype(str).apply(
                                lambda x: 1 if x.strip() and x.lower() != "nan" else 0
                            ).sum()
                        )
            ada_peta = os.path.exists(os.path.join(FOLDER_PETA, "peta_1.jpg")) and os.path.exists(
                os.path.join(FOLDER_PETA, "peta_2.jpg")
            )
            ada_ttd = any(
                os.path.exists(os.path.join(FOLDER_TTD, f))
                for f in ["ttd_day.png", "ttd_night.png", "ttd_diperiksa.png", "ttd_disetujui.png"]
            )

            def ic(ok):
                return "OK" if ok else "—"

            st.info(
                "Kelengkapan " + tgl_pdf.strftime("%d %B %Y") + "\n\n"
                + ic(n_ds >= total_tp) + " Day " + str(n_ds) + "/" + str(total_tp)
                + ((" kurang: " + ", ".join(belum_ds)) if belum_ds else "") + "\n"
                + ic(n_ns >= total_tp) + " Night " + str(n_ns) + "/" + str(total_tp)
                + ((" kurang: " + ", ".join(belum_ns)) if belum_ns else "") + "\n"
                + ic(ada_kes) + " Kesimpulan\n"
                + ic(jml_foto > 0) + " Foto: " + str(jml_foto) + "\n"
                + ic(ada_peta) + " Peta\n"
                + ic(ada_ttd) + " Tanda tangan"
            )
            if role in ["petugas", "admin"] and status_now in ["Draft", "Ditolak"]:
                list_email_dip = load_email_list(FILE_EMAIL_DIPERIKSA, ["dltdelta84@gmail.com"])
                email_tujuan = st.selectbox(
                    "Email notifikasi Environment Officer",
                    options=list_email_dip if list_email_dip else ["(isi Kelola Notifikasi)"],
                    key="email_ajukan",
                )
                if st.button("Ajukan Laporan untuk Diperiksa", type="primary", key="btn_ajukan"):
                    if not list_email_dip:
                        st.error("Belum ada email Diperiksa.")
                    else:
                        upsert_status(
                            tgl_pdf_str, Status="Menunggu Diperiksa",
                            Diajukan_Oleh=str(nama_user),
                            Waktu_Ajukan=datetime.now().strftime("%Y-%m-%d %H:%M"),
                            Email_Diperiksa=str(email_tujuan),
                            Catatan_Tolak="", Ditolak_Oleh="", Waktu_Tolak="",
                        )
                        body = (
                            "Yth. Bapak/Ibu Environment Officer,\n\n"
                            "Dengan hormat, kami sampaikan bahwa laporan pemantauan udara sesaat tanggal "
                            + tgl_pdf_str + " telah diajukan oleh " + str(nama_user)
                            + " dan menunggu pemeriksaan Bapak/Ibu.\n\n"
                            "Status saat ini: Menunggu Diperiksa."
                            + teks_email_footer()
                        )
                        ok, msg = kirim_email(
                            email_tujuan,
                            "[PEPC Zona 12] Laporan " + tgl_pdf_str + " menunggu diperiksa",
                            body,
                        )
                        st.success("Diajukan. " + msg) if ok else st.warning("Diajukan. Email: " + msg)
                        st.rerun()
            elif role in ["petugas", "admin"] and status_now == "Menunggu Diperiksa":
                st.info("Menunggu Environment Officer.")
            elif role in ["petugas", "admin"] and status_now == "Menunggu Disetujui":
                st.info("Menunggu Superintendent HSSE.")
            boleh_pdf = status_now == "Final" or role in ["diperiksa", "disetujui", "admin"]
            if boleh_pdf:
                label = "Generate & Download PDF" if status_now == "Final" else "Preview PDF (review)"
                if st.button(label, type="primary", key="btn_pdf_gen"):
                    with st.spinner("Membuat PDF..."):
                        try:
                            pdf_data = buat_pdf_laporan(df_db, tgl_pdf)
                            st.download_button(
                                "Download PDF", data=pdf_data,
                                file_name="Laporan_Pemantauan_Udara_" + tgl_pdf.strftime("%Y%m%d") + ".pdf",
                                mime="application/pdf",
                            )
                            st.success("PDF siap.")
                        except Exception as e:
                            st.error("Gagal: " + str(e))
            else:
                st.warning("PDF petugas hanya setelah Final.")

    elif pilihan_menu == "3b. Persetujuan Laporan":
        st.markdown("### Persetujuan Laporan")
        if role not in ["diperiksa", "disetujui", "admin"]:
            st.warning("Akses ditolak.")
            return
        df_st = load_status_laporan()
        if role in ["diperiksa", "admin"]:
            st.markdown("#### Menunggu Diperiksa")
            pending1 = df_st[df_st["Status"] == "Menunggu Diperiksa"] if not df_st.empty else pd.DataFrame()
            if pending1.empty:
                st.info("Tidak ada antrian.")
            else:
                kol = [c for c in ["Tanggal", "Diajukan_Oleh", "Waktu_Ajukan", "Email_Diperiksa", "Status"] if c in pending1.columns]
                st.dataframe(pending1[kol], use_container_width=True, hide_index=True)
                tgl_pilih = st.selectbox("Pilih tanggal", pending1["Tanggal"].astype(str).tolist(), key="tgl_periksa")
                list_email_dis = load_email_list(FILE_EMAIL_DISETUJUI, ["delta.unigoro@gmail.com"])
                email_sup = st.selectbox(
                    "Email Superintendent",
                    options=list_email_dis if list_email_dis else ["(isi Kelola Notifikasi)"],
                    key="email_ke_sup",
                )
                catatan1 = st.text_area("Catatan (wajib jika tolak)", key="cat_periksa")
                c1, c2 = st.columns(2)
                with c1:
                    if st.button("Setujui - Kirim ke Superintendent", type="primary", key="ok_periksa"):
                        if not list_email_dis:
                            st.error("Belum ada email Disetujui.")
                        else:
                            upsert_status(
                                str(tgl_pilih), Status="Menunggu Disetujui",
                                Diperiksa_Oleh=str(nama_user),
                                Waktu_Diperiksa=datetime.now().strftime("%Y-%m-%d %H:%M"),
                                Email_Disetujui=str(email_sup), Catatan_Tolak="",
                            )
                            body = (
                                "Yth. Bapak/Ibu Superintendent Field HSSE,\n\n"
                                "Dengan hormat, laporan pemantauan udara sesaat tanggal "
                                + str(tgl_pilih) + " telah diperiksa oleh " + str(nama_user)
                                + " dan menunggu persetujuan final Bapak/Ibu.\n\n"
                                "Status saat ini: Menunggu Disetujui."
                                + teks_email_footer()
                            )
                            ok, msg = kirim_email(
                                email_sup,
                                "[PEPC Zona 12] Laporan " + str(tgl_pilih) + " menunggu disetujui",
                                body,
                            )
                            st.success("OK. " + msg) if ok else st.warning("Status OK. Email: " + msg)
                            st.rerun()
                with c2:
                    if st.button("Tolak", key="tolak_periksa"):
                        if not catatan1.strip():
                            st.error("Catatan wajib.")
                        else:
                            upsert_status(
                                str(tgl_pilih), Status="Ditolak",
                                Ditolak_Oleh=str(nama_user),
                                Waktu_Tolak=datetime.now().strftime("%Y-%m-%d %H:%M"),
                                Catatan_Tolak=str(catatan1.strip()),
                            )
                            st.warning("Ditolak.")
                            st.rerun()
        st.markdown("---")
        if role in ["disetujui", "admin"]:
            st.markdown("#### Menunggu Disetujui")
            pending2 = df_st[df_st["Status"] == "Menunggu Disetujui"] if not df_st.empty else pd.DataFrame()
            if pending2.empty:
                st.info("Tidak ada antrian.")
            else:
                kol2 = [c for c in ["Tanggal", "Diajukan_Oleh", "Diperiksa_Oleh", "Waktu_Diperiksa", "Status"] if c in pending2.columns]
                st.dataframe(pending2[kol2], use_container_width=True, hide_index=True)
                tgl_pilih2 = st.selectbox("Pilih tanggal", pending2["Tanggal"].astype(str).tolist(), key="tgl_setujui")
                catatan2 = st.text_area("Catatan (wajib jika tolak)", key="cat_setujui")
                c3, c4 = st.columns(2)
                with c3:
                    if st.button("Setujui Final", type="primary", key="ok_setujui"):
                        row_info = pending2[pending2["Tanggal"].astype(str) == str(tgl_pilih2)]
                        diajukan = str(row_info.iloc[-1].get("Diajukan_Oleh", "") or "") if len(row_info) > 0 else ""
                        upsert_status(
                            str(tgl_pilih2), Status="Final",
                            Disetujui_Oleh=str(nama_user),
                            Waktu_Disetujui=datetime.now().strftime("%Y-%m-%d %H:%M"),
                            Catatan_Tolak="",
                        )
                        smtp_em, _ = load_smtp_config()
                        tujuan = smtp_em if smtp_em and "@" in smtp_em else "pengirimaplikasi@gmail.com"
                        body_final = (
                            "Yth. Bapak/Ibu Petugas / Tim Monitoring,\n\n"
                            "Dengan hormat, kami informasikan bahwa laporan pemantauan udara sesaat tanggal "
                            + str(tgl_pilih2) + " telah disetujui secara final oleh " + str(nama_user)
                            + ".\nDiajukan oleh: " + (diajukan or "-")
                            + "\nStatus: Final.\n\n"
                            "Petugas dapat mengunduh PDF laporan resmi melalui aplikasi."
                            + teks_email_footer()
                        )
                        ok_mail, msg_mail = kirim_email(
                            tujuan,
                            "[PEPC Zona 12] Laporan " + str(tgl_pilih2) + " SUDAH FINAL",
                            body_final,
                        )
                        st.success("Final. " + msg_mail) if ok_mail else st.success("Final. Email: " + msg_mail)
                        st.rerun()
                with c4:
                    if st.button("Tolak", key="tolak_setujui"):
                        if not catatan2.strip():
                            st.error("Catatan wajib.")
                        else:
                            upsert_status(
                                str(tgl_pilih2), Status="Ditolak",
                                Ditolak_Oleh=str(nama_user),
                                Waktu_Tolak=datetime.now().strftime("%Y-%m-%d %H:%M"),
                                Catatan_Tolak=str(catatan2.strip()),
                            )
                            st.warning("Ditolak.")
                            st.rerun()

    elif pilihan_menu == "4. Cek Kesiapan Alat":
        st.markdown("### 4. Cek Kesiapan Alat")
        if os.path.exists(FILE_ALAT):
            df_alat = pd.read_excel(FILE_ALAT)
        else:
            df_alat = pd.DataFrame(columns=[
                "Nama_Alat", "Tanggal_Inspeksi", "Tanggal_Berikutnya",
                "Indikator", "Status", "Petugas", "Keterangan",
            ])
        hari_ini = datetime.now().date()
        if not df_alat.empty and "Tanggal_Berikutnya" in df_alat.columns:
            perlu = False
            for idx, row in df_alat.iterrows():
                try:
                    if pd.to_datetime(row["Tanggal_Berikutnya"]).date() < hari_ini and row["Status"] == "Siap Pakai":
                        df_alat.at[idx, "Status"] = "Belum dicek"
                        perlu = True
                except Exception:
                    pass
            if perlu:
                df_alat.to_excel(FILE_ALAT, index=False)
        st.markdown("#### Status Terkini")
        cols = st.columns(min(4, max(len(daftar_alat), 1)))
        for i, nama in enumerate(daftar_alat):
            with cols[i % 4]:
                last = df_alat[df_alat["Nama_Alat"] == nama].tail(1)
                if len(last) > 0:
                    st_ = last.iloc[0]["Status"]
                    if st_ == "Siap Pakai":
                        st.success("**" + nama + "** — Siap")
                    elif st_ == "Tidak Siap":
                        st.warning("**" + nama + "** — Tidak Siap")
                    else:
                        st.error("**" + nama + "** — Belum dicek")
                else:
                    st.warning("**" + nama + "** — Belum dicek")
        st.markdown("---")
        alat = st.selectbox("Pilih Alat *", daftar_alat, key="pilih_alat_cek")
        nl = (alat or "").lower()
        jenis = "SO2" if "so2" in nl else ("Noise" if any(x in nl for x in ["noise", "sound", "kebisingan"]) else "Multiple")
        with st.form("form_cek_kesiapan"):
            tgl_insp = st.date_input("Tanggal Cek *", datetime.now())
            interval = st.number_input("Interval (hari) *", 1, 365, 30)
            petugas = st.text_input("Petugas *", value=nama_user)
            if jenis == "Multiple":
                a, b, c, d = st.columns(4)
                with a:
                    chk_h2s = st.checkbox("H2S OK", value=True)
                with b:
                    chk_co = st.checkbox("CO OK", value=True)
                with c:
                    chk_o2 = st.checkbox("O2 OK", value=True)
                with d:
                    chk_lel = st.checkbox("LEL OK", value=True)
                semua_ok = chk_h2s and chk_co and chk_o2 and chk_lel
                ind = "H2S:" + ("Ya" if chk_h2s else "Tidak") + ", CO:" + ("Ya" if chk_co else "Tidak") + ", O2:" + ("Ya" if chk_o2 else "Tidak") + ", LEL:" + ("Ya" if chk_lel else "Tidak")
            elif jenis == "SO2":
                chk = st.checkbox("SO2 OK", value=True)
                semua_ok, ind = chk, "SO2:" + ("Ya" if chk else "Tidak")
            else:
                chk = st.checkbox("Noise OK", value=True)
                semua_ok, ind = chk, "Noise:" + ("Ya" if chk else "Tidak")
            ket = st.text_area("Keterangan" + (" * wajib" if not semua_ok else ""))
            if st.form_submit_button("Simpan", type="primary"):
                if not petugas.strip():
                    st.error("Petugas wajib")
                elif not semua_ok and not ket.strip():
                    st.error("Keterangan wajib")
                else:
                    status = "Siap Pakai" if semua_ok else "Tidak Siap"
                    tgl_next = (tgl_insp + timedelta(days=interval)).strftime("%Y-%m-%d")
                    baru = pd.DataFrame([{
                        "Nama_Alat": alat,
                        "Tanggal_Inspeksi": tgl_insp.strftime("%Y-%m-%d"),
                        "Tanggal_Berikutnya": tgl_next,
                        "Indikator": ind, "Status": status,
                        "Petugas": petugas.strip(),
                        "Keterangan": ket.strip() if ket else "",
                    }])
                    for col in baru.columns:
                        if col not in df_alat.columns:
                            df_alat[col] = ""
                    pd.concat([df_alat, baru], ignore_index=True).to_excel(FILE_ALAT, index=False)
                    st.success("Tersimpan: " + status)
                    st.rerun()
        if not df_alat.empty:
            st.markdown("#### Riwayat")
            st.dataframe(df_alat.sort_values("Tanggal_Inspeksi", ascending=False), use_container_width=True, hide_index=True)

    elif pilihan_menu == "5. Kelola Master Data":
        st.markdown("### 5. Kelola Master Data")
        tab1, tab2, tab3, tab4, tab5, tab6, tab7 = st.tabs([
            "Kelola Titik Pantau", "Kelola Daftar Alat", "Upload Peta Laporan",
            "Nama Pejabat", "Upload Tanda Tangan", "Kelola Kesimpulan", "Kelola Notifikasi",
        ])

        with tab1:
            st.markdown("#### Titik Pantau")
            df_titik = load_titik_df()
            st.dataframe(df_titik.drop(columns=["id"], errors="ignore"), use_container_width=True, hide_index=True)
            mode_titik = st.radio("Aksi", ["Tambah Titik Baru", "Update Titik", "Hapus Titik"], horizontal=True, key="mode_titik")
            with st.form("form_titik"):
                if mode_titik == "Tambah Titik Baru":
                    kode = st.text_input("Kode *")
                    lokasi = st.text_input("Lokasi *")
                    lat = st.text_input("Latitude *")
                    long = st.text_input("Longitude *")
                    lat_ns = st.text_input("Latitude NS")
                    long_ns = st.text_input("Longitude NS")
                elif mode_titik == "Update Titik":
                    list_kode = df_titik["Kode"].tolist() if "Kode" in df_titik.columns and not df_titik.empty else []
                    kode = st.selectbox("Kode *", list_kode)
                    row = df_titik[df_titik["Kode"] == kode].iloc[0] if kode and len(df_titik) else None
                    lokasi = st.text_input("Lokasi *", value=str(row["Lokasi"]) if row is not None else "")
                    lat = st.text_input("Latitude *", value=str(row["Latitude"]) if row is not None else "")
                    long = st.text_input("Longitude *", value=str(row["Longitude"]) if row is not None else "")
                    lat_ns = st.text_input("Latitude NS", value=str(row["Latitude_NS"]) if row is not None else "")
                    long_ns = st.text_input("Longitude NS", value=str(row["Longitude_NS"]) if row is not None else "")
                else:
                    list_kode = df_titik["Kode"].tolist() if "Kode" in df_titik.columns and not df_titik.empty else []
                    kode = st.selectbox("Hapus kode *", list_kode)
                    lokasi = lat = long = lat_ns = long_ns = ""
                if st.form_submit_button("Simpan", type="primary"):
                    try:
                        sb = get_supabase()
                        if mode_titik == "Tambah Titik Baru":
                            if not all([str(kode).strip(), lokasi.strip(), lat.strip(), long.strip()]):
                                st.error("Field wajib belum lengkap")
                            else:
                                sb.table("titik_pantau").insert({
                                    "kode": str(kode).strip(), "lokasi": lokasi.strip(),
                                    "latitude": lat.strip(), "longitude": long.strip(),
                                    "latitude_ns": lat_ns.strip() or lat.strip(),
                                    "longitude_ns": long_ns.strip() or long.strip(),
                                }).execute()
                                st.success("Ditambahkan ke Supabase")
                                st.rerun()
                        elif mode_titik == "Update Titik":
                            sb.table("titik_pantau").update({
                                "lokasi": lokasi.strip(), "latitude": lat.strip(), "longitude": long.strip(),
                                "latitude_ns": lat_ns.strip() or lat.strip(),
                                "longitude_ns": long_ns.strip() or long.strip(),
                            }).eq("kode", kode).execute()
                            st.success("Diperbarui")
                            st.rerun()
                        else:
                            sb.table("titik_pantau").delete().eq("kode", kode).execute()
                            st.success("Dihapus")
                            st.rerun()
                    except Exception as e:
                        st.error(str(e))

        with tab2:
            st.markdown("#### Daftar Alat")
            alat_list = load_daftar_alat()
            st.dataframe(pd.DataFrame({"Nama_Alat": alat_list}), use_container_width=True, hide_index=True)
            mode_alat = st.radio("Aksi", ["Tambah Alat Baru", "Update Nama Alat", "Hapus Alat"], horizontal=True, key="mode_alat")
            with st.form("form_alat"):
                if mode_alat == "Tambah Alat Baru":
                    nama_alat = st.text_input("Nama Alat *")
                    nama_lama = None
                elif mode_alat == "Update Nama Alat":
                    nama_lama = st.selectbox("Pilih alat", alat_list)
                    nama_alat = st.text_input("Nama baru *", value=nama_lama or "")
                else:
                    nama_alat = st.selectbox("Hapus", alat_list)
                    nama_lama = None
                if st.form_submit_button("Simpan", type="primary"):
                    try:
                        sb = get_supabase()
                        if mode_alat == "Tambah Alat Baru" and nama_alat.strip():
                            sb.table("daftar_alat").insert({"nama_alat": nama_alat.strip()}).execute()
                            st.success("OK")
                            st.rerun()
                        elif mode_alat == "Update Nama Alat" and nama_lama:
                            sb.table("daftar_alat").update({"nama_alat": nama_alat.strip()}).eq("nama_alat", nama_lama).execute()
                            st.success("OK")
                            st.rerun()
                        elif mode_alat == "Hapus Alat":
                            sb.table("daftar_alat").delete().eq("nama_alat", nama_alat).execute()
                            st.success("OK")
                            st.rerun()
                    except Exception as e:
                        st.error(str(e))

        with tab3:
            st.markdown("#### Upload Peta")
            os.makedirs(FOLDER_PETA, exist_ok=True)
            c1, c2 = st.columns(2)
            with c1:
                p1 = st.file_uploader("Peta 1", type=["jpg", "jpeg", "png"], key="peta1")
                path_p1 = os.path.join(FOLDER_PETA, "peta_1.jpg")
                if p1:
                    with open(path_p1, "wb") as f:
                        f.write(p1.getbuffer())
                    st.image(p1, width=250)
                    st.success("Peta 1 OK")
                elif os.path.exists(path_p1):
                    st.image(path_p1, width=250)
            with c2:
                p2 = st.file_uploader("Peta 2", type=["jpg", "jpeg", "png"], key="peta2")
                path_p2 = os.path.join(FOLDER_PETA, "peta_2.jpg")
                if p2:
                    with open(path_p2, "wb") as f:
                        f.write(p2.getbuffer())
                    st.image(p2, width=250)
                    st.success("Peta 2 OK")
                elif os.path.exists(path_p2):
                    st.image(path_p2, width=250)

        with tab4:
            st.markdown("#### Nama Pejabat PDF")
            peng = load_pengaturan_laporan()
            with st.form("form_pejabat"):
                a, b = st.columns(2)
                with a:
                    nd = st.text_input("Nama Diperiksa", value=peng["Nama_Diperiksa"])
                    jd = st.text_input("Jabatan Diperiksa", value=peng["Jabatan_Diperiksa"])
                with b:
                    ns = st.text_input("Nama Disetujui", value=peng["Nama_Disetujui"])
                    js = st.text_input("Jabatan Disetujui", value=peng["Jabatan_Disetujui"])
                if st.form_submit_button("Simpan", type="primary"):
                    try:
                        sb = get_supabase()
                        payload = {
                            "nama_diperiksa": nd, "jabatan_diperiksa": jd,
                            "nama_disetujui": ns, "jabatan_disetujui": js,
                        }
                        cek = sb.table("pengaturan_laporan").select("id").limit(1).execute()
                        if cek.data:
                            sb.table("pengaturan_laporan").update(payload).eq("id", cek.data[0]["id"]).execute()
                        else:
                            sb.table("pengaturan_laporan").insert(payload).execute()
                        st.success("Disimpan ke Supabase")
                        st.rerun()
                    except Exception as e:
                        st.error(str(e))

        with tab5:
            st.markdown("#### Upload Tanda Tangan")
            os.makedirs(FOLDER_TTD, exist_ok=True)
            for label, fname, key in [
                ("Day Shift", "ttd_day.png", "ttd_ds"),
                ("Night Shift", "ttd_night.png", "ttd_ns"),
                ("Diperiksa", "ttd_diperiksa.png", "ttd_dip"),
                ("Disetujui", "ttd_disetujui.png", "ttd_dis"),
            ]:
                st.markdown("**" + label + "**")
                f = st.file_uploader("Upload " + label, type=["png", "jpg", "jpeg"], key=key)
                path = os.path.join(FOLDER_TTD, fname)
                if f:
                    with open(path, "wb") as out:
                        out.write(f.getbuffer())
                    st.image(f, width=180)
                    st.success("OK")
                elif os.path.exists(path):
                    st.image(path, width=180)

        with tab6:
            st.markdown("#### Kelola Kesimpulan")
            lst = load_kesimpulan_list()
            st.dataframe(pd.DataFrame({"Kesimpulan": lst}), use_container_width=True, hide_index=True)
            mode_kes = st.radio("Aksi", ["Tambah", "Update", "Hapus"], horizontal=True, key="mode_kes")
            with st.form("form_kes"):
                if mode_kes == "Tambah":
                    teks = st.text_area("Teks *", height=80)
                    idx = None
                elif mode_kes == "Update":
                    idx = st.selectbox("Pilih", range(len(lst)), format_func=lambda i: lst[i])
                    teks = st.text_area("Teks baru *", value=lst[idx] if lst else "", height=80)
                else:
                    idx = st.selectbox("Hapus", range(len(lst)), format_func=lambda i: lst[i])
                    teks = ""
                if st.form_submit_button("Simpan", type="primary"):
                    try:
                        sb = get_supabase()
                        if mode_kes == "Tambah" and teks.strip():
                            sb.table("kesimpulan").insert({"teks": teks.strip()}).execute()
                            st.rerun()
                        elif mode_kes == "Update" and teks.strip() and lst:
                            old = lst[idx]
                            sb.table("kesimpulan").update({"teks": teks.strip()}).eq("teks", old).execute()
                            st.rerun()
                        elif mode_kes == "Hapus" and len(lst) > 1:
                            sb.table("kesimpulan").delete().eq("teks", lst[idx]).execute()
                            st.rerun()
                    except Exception as e:
                        st.error(str(e))

        with tab7:
            st.markdown("#### Kelola Notifikasi Email")
            try:
                _ = st.secrets["smtp"]["email"]
                st.info("SMTP aktif dari Streamlit Secrets.")
            except Exception:
                sm_em, sm_pw = load_smtp_config()
                with st.form("form_smtp"):
                    smtp_email = st.text_input("Email pengirim", value=sm_em)
                    smtp_pass = st.text_input("App Password", value=sm_pw, type="password")
                    if st.form_submit_button("Simpan SMTP lokal", type="primary"):
                        save_smtp_config(smtp_email, smtp_pass)
                        st.success("Disimpan")
                        st.rerun()
            emails_dip = load_email_list(FILE_EMAIL_DIPERIKSA, ["dltdelta84@gmail.com"])
            st.markdown("##### Email Diperiksa")
            st.dataframe(pd.DataFrame({"Email": emails_dip}), use_container_width=True, hide_index=True)
            mode_ed = st.radio("Aksi Diperiksa", ["Tambah", "Hapus"], horizontal=True, key="mode_ed")
            with st.form("form_email_dip"):
                em = st.text_input("Email") if mode_ed == "Tambah" else st.selectbox("Hapus", emails_dip or ["-"])
                if st.form_submit_button("Simpan Diperiksa"):
                    try:
                        sb = get_supabase()
                        if mode_ed == "Tambah" and em.strip() and "@" in em:
                            sb.table("email_diperiksa").insert({"email": em.strip()}).execute()
                            st.rerun()
                        elif mode_ed == "Hapus" and em in emails_dip:
                            sb.table("email_diperiksa").delete().eq("email", em).execute()
                            st.rerun()
                    except Exception as e:
                        st.error(str(e))
            emails_dis = load_email_list(FILE_EMAIL_DISETUJUI, ["delta.unigoro@gmail.com"])
            st.markdown("##### Email Disetujui")
            st.dataframe(pd.DataFrame({"Email": emails_dis}), use_container_width=True, hide_index=True)
            mode_es = st.radio("Aksi Disetujui", ["Tambah", "Hapus"], horizontal=True, key="mode_es")
            with st.form("form_email_dis"):
                em2 = st.text_input("Email ") if mode_es == "Tambah" else st.selectbox("Hapus ", emails_dis or ["-"])
                if st.form_submit_button("Simpan Disetujui"):
                    try:
                        sb = get_supabase()
                        if mode_es == "Tambah" and em2.strip() and "@" in em2:
                            sb.table("email_disetujui").insert({"email": em2.strip()}).execute()
                            st.rerun()
                        elif mode_es == "Hapus" and em2 in emails_dis:
                            sb.table("email_disetujui").delete().eq("email", em2).execute()
                            st.rerun()
                    except Exception as e:
                        st.error(str(e))
            if st.button("Tes kirim email Diperiksa"):
                tujuan = emails_dip[0] if emails_dip else "dltdelta84@gmail.com"
                ok, msg = kirim_email(
                    tujuan,
                    "[PEPC Zona 12] Tes notifikasi",
                    "Yth. Bapak/Ibu,\n\nIni adalah email uji sistem pemantauan udara."
                    + teks_email_footer(),
                )
                st.success(msg) if ok else st.error(msg)

    elif pilihan_menu == "6. Kelola User":
        st.markdown("### 6. Kelola User")
        st.info("Login memakai tabel Supabase **users**.")
        try:
            res = get_supabase().table("users").select("id, username, nama, role").execute()
            df_u = pd.DataFrame(res.data) if res.data else pd.DataFrame()
            if not df_u.empty:
                st.dataframe(df_u, use_container_width=True, hide_index=True)
            else:
                st.info("Belum ada user.")
        except Exception as e:
            st.error(str(e))